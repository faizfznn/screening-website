import openpyxl

file_path = r'd:/AKU File/BEM/Salinan dari KRITERIA & JADWAL GRAND RECRUITMENT JUNIOR STAFF SGE 2026.xlsx'
wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

with open('excel_analysis.txt', 'w', encoding='utf-8') as f:
    f.write("=== DAFTAR SEMUA SHEET ===\n")
    for name in wb.sheetnames:
        f.write(f"- {name}\n")
    
    # 1. Inspect ALL Registrants
    f.write("\n=== 1. ALL REGISTRANTS HEADER & SAMPLE ===\n")
    for sname in wb.sheetnames:
        if 'registrant' in sname.lower():
            ws = wb[sname]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx < 10:
                    f.write(f"Row {r_idx+1}: {row[:12]}\n")
            break

    # 2. Inspect Daily Sheets
    f.write("\n=== 2. DAILY SHEETS SAMPLE (05-10-2026, 08-10-2026, etc) ===\n")
    for sname in wb.sheetnames:
        if any(d in sname for d in ['05-10', '08-10', '09-10']):
            ws = wb[sname]
            f.write(f"\n--- Sheet: {sname} ---\n")
            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx < 8:
                    f.write(f"Row {r_idx+1}: {row[:15]}\n")
            break

    # 3. Inspect Schedule Sheets
    f.write("\n=== 3. SCHEDULE SHEETS (BoD, Staff, etc) ===\n")
    for sname in wb.sheetnames:
        if 'schedule' in sname.lower():
            ws = wb[sname]
            f.write(f"\n--- Sheet: {sname} ---\n")
            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx < 8:
                    f.write(f"Row {r_idx+1}: {row[:15]}\n")
            break

print("Analysis written to excel_analysis.txt")
