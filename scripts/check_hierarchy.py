import openpyxl

file_path = r'd:/AKU File/BEM/Salinan dari KRITERIA & JADWAL GRAND RECRUITMENT JUNIOR STAFF SGE 2026.xlsx'
wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

schedule_sheets = ['📅 BoD Schedule', '📅 C-Level Schedule', '📅 IRE Schedule', '📅 Mentor Schedule', '📅 Staff Schedule']

with open('schedule_hierarchy.txt', 'w', encoding='utf-8') as f:
    for sname in schedule_sheets:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        f.write(f"=== Sheet: {sname} ===\n")
        curr_minbur = ''
        entries = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx < 3: continue
            col_b = str(row[1] or '').strip()
            col_c = str(row[2] or '').strip()
            if 'REMARKS' in col_b.upper() or 'REMARKS' in col_c.upper():
                break
            if 'FREE' in col_b.upper() or 'CAN' in col_b.upper():
                continue
            if col_b:
                curr_minbur = col_b
            if col_c:
                entries.append((curr_minbur, col_c))
        
        grouped = {}
        for mb, name in entries:
            grouped.setdefault(mb, []).append(name)
        
        for mb, names in grouped.items():
            f.write(f"  * Minbur: {mb} ({len(names)} orang) -> {names}\n")
        f.write("\n")

print("Written to schedule_hierarchy.txt")
