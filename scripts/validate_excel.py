import openpyxl, datetime

file_path = r'd:/AKU File/BEM/Salinan dari KRITERIA & JADWAL GRAND RECRUITMENT JUNIOR STAFF SGE 2026.xlsx'
wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

with open('excel_validation_results.txt', 'w', encoding='utf-8') as f:
    # 1. Test Registrants count
    ws_reg = wb['📑 ALL Registrants']
    reg_count = 0
    sample_reg = []
    for row in ws_reg.iter_rows(values_only=True):
        no = row[0]
        if isinstance(no, (int, float)) and no > 0:
            nama = str(row[1] or '').strip()
            if nama:
                reg_count += 1
                if len(sample_reg) < 5:
                    sample_reg.append((int(no), nama, row[2], row[3], row[8]))
    
    f.write(f"1. ALL REGISTRANTS:\n- Total: {reg_count}\n- Samples: {sample_reg}\n\n")

    # 2. Test Daily Plotting merges
    daily_stats = {}
    for sname in wb.sheetnames:
        if '-' in sname and len(sname.split('-')) == 3:
            ws = wb[sname]
            count = 0
            for row in ws.iter_rows(values_only=True):
                if len(row) > 5 and isinstance(row[1], (int, float)) and row[1] > 0:
                    count += 1
            daily_stats[sname] = count
    
    f.write(f"2. DAILY PLOTTED CANDIDATES:\n")
    for sname, c in sorted(daily_stats.items()):
        f.write(f"- Sheet {sname}: {c} sesi interview\n")

    # 3. Test Schedule Availability
    schedule_sheets = ['📅 BoD Schedule', '📅 IRE Schedule', '📅 C-Level Schedule', '📅 Mentor Schedule', '📅 Staff Schedule']
    panelists_by_cat = {}
    bisa_slots = 0
    tidak_slots = 0

    for sname in schedule_sheets:
        if sname in wb.sheetnames:
            ws = wb[sname]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 4: continue
            
            date_row = rows[1]
            time_row = rows[2]
            
            col_map = {}
            active_date = ''
            for c in range(3, len(time_row)):
                raw_d = date_row[c] if c < len(date_row) else None
                if raw_d:
                    if isinstance(raw_d, datetime.datetime):
                        active_date = raw_d.strftime('%Y-%m-%d')
                    else:
                        active_date = str(raw_d).strip()
                
                time_slot = str(time_row[c] or '').strip().replace('.', ':')
                if active_date and time_slot and '-' in time_slot:
                    col_map[c] = (active_date, time_slot)
            
            curr_cat = ""
            for r in range(3, len(rows)):
                row = rows[r]
                col_b = str(row[1] or '').strip()
                col_c = str(row[2] or '').strip()
                if 'REMARKS' in col_b.upper() or 'REMARKS' in col_c.upper():
                    break
                if 'FREE' in col_b.upper() or 'CAN' in col_b.upper():
                    continue
                if col_b: curr_cat = col_b
                if not col_c: continue
                
                cat_key = curr_cat or sname.replace('📅', '').strip()
                if cat_key not in panelists_by_cat:
                    panelists_by_cat[cat_key] = []
                if col_c not in panelists_by_cat[cat_key]:
                    panelists_by_cat[cat_key].append(col_c)
                
                for c, slot in col_map.items():
                    if c < len(row):
                        val = str(row[c] or '').strip().lower()
                        if 'bisa' in val or 'free' in val or val == 'ya':
                            bisa_slots += 1
                        elif 'tidak' in val or 'can' in val or 'gabisa' in val:
                            tidak_slots += 1

    f.write(f"\n3. PANELISTS & AVAILABILITY:\n")
    f.write(f"- Total Kategori: {len(panelists_by_cat)}\n")
    for cat, p_list in panelists_by_cat.items():
        f.write(f"  * {cat} ({len(p_list)} orang): {p_list}\n")
    f.write(f"- Total Slot BISA di Excel: {bisa_slots}\n")
    f.write(f"- Total Slot TIDAK di Excel: {tidak_slots}\n")

print("Validation completed successfully.")
